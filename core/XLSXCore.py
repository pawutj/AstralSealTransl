"""
XLSX Processing Core for AstralSealTransl

Handles reading visual novel scripts from XLSX files and writing
translations back to XLSX format using JSONLine intermediate format.
"""

import json
from pathlib import Path
from typing import Dict, List, Optional, Union, Tuple
from dataclasses import dataclass

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.workbook.workbook import Workbook

from core.CConfig import CConfig, XLSXConfig


@dataclass
class RowData:
    """Single row data structure"""
    row_id: int          # Sequential ID for JSONLine (1, 2, 3, ...)
    excel_row: int       # Original Excel row number (for writing back)
    name: str
    src: str


class XLSXCore:
    """
    Core XLSX processing engine.

    Converts XLSX visual novel scripts to JSONLine format for translation,
    and writes translated JSONLine back to XLSX files.

    Usage:
        config = CConfig("config.yaml")
        xlsx = XLSXCore(config)

        jsonline = xlsx.readXlsx()
        print(jsonline)

        translated = '{"id":1,"dst":"แปลภาษาไทย"}'
        xlsx.writeXlsx(translated)
    """

    def __init__(self, config: CConfig) -> None:
        """
        Initialize XLSX processor.

        Args:
            config: Configuration instance with xlsx settings
        """
        self.config = config
        self.xlsx_config = config.xlsx
        self.row_mapping: Dict[int, int] = {}  # Maps JSONLine ID -> Excel row number

    def readXlsx(self, sheet_name: Optional[str] = None) -> str:
        """
        Read XLSX file and convert to JSONLine format.

        Args:
            sheet_name: Optional sheet name to read. If None, uses config.sheetName[0]

        Returns:
            JSONLine string with format: {"id":1,"name":"...","src":"..."}

        Raises:
            FileNotFoundError: If XLSX file doesn't exist
            ValueError: If sheet or columns don't exist
        """
        file_path = Path(self.xlsx_config.filePath)

        self._validate_file_exists(file_path)

        # Use provided sheet name or default to first configured sheet
        target_sheet = sheet_name if sheet_name else self.xlsx_config.sheetName[0]

        workbook = load_workbook(file_path, read_only=True, data_only=True)
        sheet = self._get_sheet(workbook, target_sheet)

        if self.xlsx_config.validateColumns:
            self._validate_columns(sheet)

        rows_data = self._extract_rows(sheet)
        jsonline = self._format_as_jsonline(rows_data)

        workbook.close()
        return jsonline

    def writeXlsx(self, jsonline_input: str) -> None:
        """
        Write translated JSONLine to XLSX file.

        Supports both single-target and dual-target modes:
        - Single: {"id":1,"dst":"..."}
        - Dual: {"id":1,"dst1":"direct","dst2":"localized"}

        Args:
            jsonline_input: JSONLine string with translations

        Raises:
            FileNotFoundError: If source XLSX doesn't exist
            ValueError: If JSONLine format is invalid or dual-target without targetColumn2
        """
        translations = self._parse_jsonline(jsonline_input)

        source_path = Path(self.xlsx_config.filePath)
        self._validate_file_exists(source_path)

        workbook = load_workbook(source_path)
        sheet = self._get_sheet(workbook, self.xlsx_config.sheetName)

        column_index = self._get_or_create_column_index(sheet, self.xlsx_config.targetColumn)

        # Prepare output path
        output_path = Path(self.xlsx_config.outputPath)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        # Detect dual-target mode by checking first translation value type
        if translations and isinstance(next(iter(translations.values())), tuple):
            # Dual-target mode: write to both columns
            if not self.xlsx_config.targetColumn2:
                raise ValueError(
                    "Dual-target translations detected but targetColumn2 not configured in config.yaml"
                )

            column_index2 = self._get_or_create_column_index(sheet, self.xlsx_config.targetColumn2)
            written_count = self._write_dual_translations(
                sheet, column_index, column_index2, translations
            )
            print(f"✅ Written {written_count} dual translations (direct + localized) to {output_path}")
        else:
            # Single-target mode: write to one column
            written_count = self._write_translations(sheet, column_index, translations)
            print(f"✅ Written {written_count} translations to {output_path}")

        workbook.save(output_path)
        workbook.close()

    def writeMultipleSheets(self, translations_by_sheet: Dict[str, str]) -> None:
        """
        Write translations for multiple sheets to a single XLSX file.

        This method efficiently handles multi-sheet writing by:
        1. Loading the source workbook once
        2. Writing translations to each specified sheet
        3. Saving the complete workbook once

        Args:
            translations_by_sheet: Dict mapping sheet name -> JSONLine translations
                Example: {
                    "s4_1": '{"id":1,"dst":"..."}\\n{"id":2,"dst":"..."}',
                    "s4_2": '{"id":1,"dst":"..."}\\n{"id":2,"dst":"..."}'
                }

        Raises:
            FileNotFoundError: If source XLSX doesn't exist
            ValueError: If sheet not found or JSONLine format invalid
        """
        source_path = Path(self.xlsx_config.filePath)
        self._validate_file_exists(source_path)

        # Load workbook once
        workbook = load_workbook(source_path)

        # Track statistics
        total_written = 0
        sheets_processed = []

        try:
            for sheet_name, jsonline_input in translations_by_sheet.items():
                # Parse translations
                translations = self._parse_jsonline(jsonline_input)

                # Get target sheet
                sheet = self._get_sheet(workbook, sheet_name)

                # Detect and write translations (reuse existing logic)
                column_index = self._get_or_create_column_index(sheet, self.xlsx_config.targetColumn)

                if translations and isinstance(next(iter(translations.values())), tuple):
                    # Dual-target mode
                    if not self.xlsx_config.targetColumn2:
                        raise ValueError(
                            f"Sheet '{sheet_name}': Dual-target translations detected "
                            f"but targetColumn2 not configured"
                        )
                    column_index2 = self._get_or_create_column_index(sheet, self.xlsx_config.targetColumn2)
                    written_count = self._write_dual_translations(
                        sheet, column_index, column_index2, translations
                    )
                else:
                    # Single-target mode
                    written_count = self._write_translations(sheet, column_index, translations)

                total_written += written_count
                sheets_processed.append(f"{sheet_name} ({written_count} sentences)")

            # Save workbook once (all sheets written)
            output_path = Path(self.xlsx_config.outputPath)
            output_path.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(output_path)

            print(f"✅ Written {len(translations_by_sheet)} sheets ({total_written} total sentences) to {output_path}")
            for sheet_info in sheets_processed:
                print(f"   - {sheet_info}")

        finally:
            workbook.close()

    def _validate_file_exists(self, file_path: Path) -> None:
        """Validate that file exists"""
        if not file_path.exists():
            raise FileNotFoundError(f"XLSX file not found: {file_path}")

    def _get_sheet(self, workbook: Workbook, sheet_name: str) -> Worksheet:
        """Get worksheet by name with validation"""
        if sheet_name not in workbook.sheetnames:
            available = ", ".join(workbook.sheetnames)
            raise ValueError(
                f"Sheet '{sheet_name}' not found. "
                f"Available sheets: {available}"
            )
        return workbook[sheet_name]

    def _validate_columns(self, sheet: Worksheet) -> None:
        """Validate that required columns exist in sheet"""
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = [str(h) if h else "" for h in header_row]

        required = [
            self.xlsx_config.nameColumn,
            self.xlsx_config.srcColumn
        ]

        missing = [col for col in required if col not in headers]

        if missing:
            raise ValueError(
                f"Required columns not found: {', '.join(missing)}. "
                f"Available columns: {', '.join(headers)}"
            )

    def _get_column_index(self, sheet: Worksheet, column_name: str) -> int:
        """Get column index by name (1-based)"""
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = [str(h) if h else "" for h in header_row]

        if column_name not in headers:
            raise ValueError(
                f"Column '{column_name}' not found. "
                f"Available: {', '.join(headers)}"
            )

        return headers.index(column_name) + 1

    def _get_or_create_column_index(self, sheet: Worksheet, column_name: str) -> int:
        """
        Get column index by name (1-based), or create new column if not exists.

        If column doesn't exist, adds it to the end of the header row.

        Args:
            sheet: Worksheet to check/modify
            column_name: Name of column to find or create

        Returns:
            Column index (1-based)
        """
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = [str(h) if h else "" for h in header_row]

        if column_name in headers:
            # Column exists, return its index
            return headers.index(column_name) + 1

        # Column doesn't exist - create it at the end
        new_column_index = len(headers) + 1
        sheet.cell(row=1, column=new_column_index).value = column_name
        print(f"ℹ️  Created new column '{column_name}' at position {new_column_index}")

        return new_column_index

    def _extract_rows(self, sheet: Worksheet) -> List[RowData]:
        """Extract data rows from worksheet, skipping empty narrator rows"""
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = [str(h) if h else "" for h in header_row]

        name_index = headers.index(self.xlsx_config.nameColumn)
        src_index = headers.index(self.xlsx_config.srcColumn)

        rows_data = []
        row_id = 1
        excel_row_num = 2  # Excel rows start at 2 (after header)

        for row in sheet.iter_rows(min_row=2, values_only=True):
            # Handle name_value first (might be float/int from Excel)
            name_value = row[name_index]
            if name_value is None:
                name_value = ""
            else:
                name_value = str(name_value)  # Convert to string (handles float/int)

            # Handle src_value (might be float/int from Excel)
            src_value = row[src_index]
            if src_value is None:
                src_value = ""
            else:
                src_value = str(src_value)  # Convert to string (handles float/int)

            # Skip only if BOTH conditions met:
            # 1. No name (narrator line)
            # 2. Empty source text
            # Keep rows with name even if src is empty (dialogue placeholder)
            if not name_value.strip() and not src_value.strip():
                excel_row_num += 1
                continue

            rows_data.append(RowData(
                row_id=row_id,
                excel_row=excel_row_num,
                name=str(name_value),
                src=str(src_value)
            ))

            # Store mapping for later use in writeXlsx
            self.row_mapping[row_id] = excel_row_num

            row_id += 1
            excel_row_num += 1

        return rows_data

    def _format_as_jsonline(self, rows_data: List[RowData]) -> str:
        """Convert rows to JSONLine format"""
        lines = []

        for row in rows_data:
            json_obj = {
                "id": row.row_id,
                "name": row.name,
                "src": row.src
            }
            lines.append(json.dumps(json_obj, ensure_ascii=False))

        return "\n".join(lines)

    def _parse_jsonline(self, jsonline_input: str) -> Union[Dict[int, str], Dict[int, Tuple[str, str]]]:
        """
        Parse JSONLine input to translation dictionary.

        Returns:
            - Dict[int, str]: Single-target mode {1: "translation"}
            - Dict[int, Tuple[str, str]]: Dual-target mode {1: ("direct", "localized")}

        Raises:
            ValueError: If JSON format is invalid or missing required keys
        """
        translations = {}
        is_dual_target = None

        for line_num, line in enumerate(jsonline_input.strip().split('\n'), start=1):
            if not line.strip():
                continue

            try:
                data = json.loads(line)
            except json.JSONDecodeError as e:
                raise ValueError(
                    f"Invalid JSON at line {line_num}: {line}\n"
                    f"Error: {e}"
                )

            # Validate required fields
            if 'id' not in data:
                raise ValueError(f"Missing 'id' at line {line_num}: {line}")

            row_id = data['id']
            if not isinstance(row_id, int):
                raise ValueError(f"Invalid 'id' type at line {line_num}: expected int")

            # Detect mode and validate consistency
            has_dual = 'dst1' in data and 'dst2' in data
            has_single = 'dst' in data

            if is_dual_target is None:
                # First line determines mode
                is_dual_target = has_dual
            elif is_dual_target != has_dual:
                raise ValueError(
                    f"Inconsistent output format at line {line_num}: "
                    f"expected {'dual' if is_dual_target else 'single'}-target"
                )

            # Extract translations
            if is_dual_target:
                if not has_dual:
                    raise ValueError(f"Missing dst1/dst2 at line {line_num}: {line}")
                translations[row_id] = (str(data['dst1']), str(data['dst2']))
            else:
                if not has_single:
                    raise ValueError(f"Missing dst at line {line_num}: {line}")
                translations[row_id] = str(data['dst'])

        return translations

    def _write_translations(
        self,
        sheet: Worksheet,
        column_index: int,
        translations: Dict[int, str]
    ) -> int:
        """
        Write translations to worksheet using row mapping.

        Returns:
            Number of rows written
        """
        written_count = 0

        for json_id, translated_text in translations.items():
            # Use mapping to find correct Excel row
            if json_id not in self.row_mapping:
                print(f"⚠️  Warning: JSONLine ID {json_id} not found in row mapping, skipping")
                continue

            excel_row = self.row_mapping[json_id]
            sheet.cell(row=excel_row, column=column_index).value = translated_text
            written_count += 1

        return written_count

    def _write_dual_translations(
        self,
        sheet: Worksheet,
        column_index1: int,
        column_index2: int,
        translations: Dict[int, Tuple[str, str]]
    ) -> int:
        """
        Write dual translations (direct + localized) to two columns.

        Returns:
            Number of rows written
        """
        written_count = 0

        for json_id, (direct_trans, localized_trans) in translations.items():
            # Use mapping to find correct Excel row
            if json_id not in self.row_mapping:
                print(f"⚠️  Warning: JSONLine ID {json_id} not found in row mapping, skipping")
                continue

            excel_row = self.row_mapping[json_id]

            # Write direct translation to column 1
            sheet.cell(row=excel_row, column=column_index1).value = direct_trans

            # Write localized translation to column 2
            sheet.cell(row=excel_row, column=column_index2).value = localized_trans

            written_count += 1

        return written_count
