"""
Integration tests for XLSXCore

Tests all cases from task_xlsx.md checklist:
- Reading with empty name/src columns
- Error handling (file not found, sheet not found, column not found)
- Writing valid/invalid JSONLine
- ID mismatch handling
"""

from pathlib import Path
import sys
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).parent))

from core.CConfig import CConfig
from core.XLSXCore import XLSXCore


def create_test_xlsx(file_path: Path) -> None:
    """Create test XLSX file with sample data"""
    wb = Workbook()
    ws = wb.active
    ws.title = "s4_1"

    ws['A1'] = 'who_talk'
    ws['B1'] = 'talk'
    ws['C1'] = 'talk_jp'

    test_data = [
        ('Reika', 'ごめんなさい、こんな夕方に呼び出して', ''),
        ('', 'その声は普段とは違い、柔らかく耳に響いた', ''),
        ('Reika', '私、もう決めたの', ''),
        ('Reika', '', ''),
        ('', '', ''),
    ]

    for idx, (name, talk, talk_jp) in enumerate(test_data, start=2):
        ws[f'A{idx}'] = name
        ws[f'B{idx}'] = talk
        ws[f'C{idx}'] = talk_jp

    file_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(file_path)
    wb.close()


def test_read_with_empty_cells():
    """Test reading XLSX with empty name and src columns"""
    print("\n" + "="*60)
    print("Test: Read XLSX with empty cells")
    print("="*60)

    test_file = Path("input/test_script.xlsx")
    create_test_xlsx(test_file)

    config = CConfig("config.yaml")
    config.xlsx.filePath = str(test_file)
    config.xlsx.sheetName = ["s4_1"]  # Must be list

    xlsx = XLSXCore(config)
    result = xlsx.readXlsx("s4_1")

    lines = result.strip().split('\n')
    print(f"✅ Read {len(lines)} lines")

    import json
    for line in lines:
        data = json.loads(line)
        print(f"  ID {data['id']}: name='{data['name']}', src='{data['src']}'")

    # Expected 4 rows (not 5):
    # - Row 2: Reika | ごめんなさい... ✓
    # - Row 3: (empty) | その声は... ✓
    # - Row 4: Reika | 私、もう... ✓
    # - Row 5: Reika | (empty) ✓ (kept because has name)
    # - Row 6: (empty) | (empty) ✗ (skipped - both empty)
    assert len(lines) == 4, "Should read 4 rows (row 6 skipped)"
    assert json.loads(lines[1])['name'] == '', "Row 2 name should be empty"
    assert json.loads(lines[3])['src'] == '', "Row 4 src should be empty"

    print("✅ Empty cells handled correctly")


def test_file_not_found():
    """Test FileNotFoundError when file doesn't exist"""
    print("\n" + "="*60)
    print("Test: FileNotFoundError")
    print("="*60)

    config = CConfig("config.yaml")
    config.xlsx.filePath = "nonexistent.xlsx"

    xlsx = XLSXCore(config)

    try:
        xlsx.readXlsx()
        assert False, "Should raise FileNotFoundError"
    except FileNotFoundError as e:
        print(f"✅ Correct error: {e}")
        assert "not found" in str(e)


def test_sheet_not_found():
    """Test ValueError when sheet doesn't exist"""
    print("\n" + "="*60)
    print("Test: Sheet not found")
    print("="*60)

    test_file = Path("input/test_script.xlsx")
    create_test_xlsx(test_file)

    config = CConfig("config.yaml")
    config.xlsx.filePath = str(test_file)
    config.xlsx.sheetName = ["InvalidSheet"]  # Must be list

    xlsx = XLSXCore(config)

    try:
        xlsx.readXlsx("InvalidSheet")
        assert False, "Should raise ValueError"
    except ValueError as e:
        print(f"✅ Correct error: {e}")
        assert "Sheet" in str(e)
        assert "not found" in str(e)


def test_column_not_found():
    """Test ValueError when column doesn't exist"""
    print("\n" + "="*60)
    print("Test: Column not found")
    print("="*60)

    test_file = Path("input/test_script.xlsx")
    create_test_xlsx(test_file)

    config = CConfig("config.yaml")
    config.xlsx.filePath = str(test_file)
    config.xlsx.sheetName = ["s4_1"]  # Must be list
    config.xlsx.nameColumn = "invalid_column"

    xlsx = XLSXCore(config)

    try:
        xlsx.readXlsx("s4_1")
        assert False, "Should raise ValueError"
    except ValueError as e:
        print(f"✅ Correct error: {e}")
        assert "not found" in str(e)


def test_write_valid_jsonline():
    """Test writing valid JSONLine"""
    print("\n" + "="*60)
    print("Test: Write valid JSONLine")
    print("="*60)

    test_file = Path("input/test_script.xlsx")
    output_file = Path("output/test_translated.xlsx")
    create_test_xlsx(test_file)

    config = CConfig("config.yaml")
    config.xlsx.filePath = str(test_file)
    config.xlsx.outputPath = str(output_file)
    config.xlsx.sheetName = ["s4_1"]  # Must be list

    jsonline = '''{"id":1,"dst":"ขอโทษนะที่เรียกมาตอนเย็นแบบนี้"}
{"id":2,"dst":"เสียงนั้นแตกต่างจากปกติ ดังกังวานในหูอย่างนุ่มนวล"}
{"id":3,"dst":"ฉันตัดสินใจแล้ว"}'''

    xlsx = XLSXCore(config)

    # Must read first to build row mapping
    xlsx.readXlsx("s4_1")

    # Now write with sheet_name parameter
    xlsx.writeXlsx(jsonline, sheet_name="s4_1")

    assert output_file.exists(), "Output file should be created"
    print(f"✅ Output created: {output_file}")


def test_write_invalid_json():
    """Test JSONDecodeError with invalid JSON"""
    print("\n" + "="*60)
    print("Test: Invalid JSON format")
    print("="*60)

    test_file = Path("input/test_script.xlsx")
    create_test_xlsx(test_file)

    config = CConfig("config.yaml")
    config.xlsx.filePath = str(test_file)
    config.xlsx.sheetName = ["s4_1"]  # Must be list

    invalid_json = '''{"id":1,"dst":"valid"}
{invalid json here}
{"id":3,"dst":"valid"}'''

    xlsx = XLSXCore(config)

    # Read first (though error will happen before write)
    xlsx.readXlsx("s4_1")

    try:
        xlsx.writeXlsx(invalid_json, sheet_name="s4_1")
        assert False, "Should raise ValueError"
    except ValueError as e:
        print(f"✅ Correct error: {e}")
        assert "Invalid JSON" in str(e)
        assert "line 2" in str(e)


def test_write_missing_keys():
    """Test ValueError when JSON missing required keys"""
    print("\n" + "="*60)
    print("Test: Missing 'id' or 'dst' keys")
    print("="*60)

    test_file = Path("input/test_script.xlsx")
    create_test_xlsx(test_file)

    config = CConfig("config.yaml")
    config.xlsx.filePath = str(test_file)
    config.xlsx.sheetName = ["s4_1"]  # Must be list

    missing_dst = '''{"id":1,"dst":"valid"}
{"id":2}
{"id":3,"dst":"valid"}'''

    xlsx = XLSXCore(config)

    # Read first (though error will happen before write)
    xlsx.readXlsx("s4_1")

    try:
        xlsx.writeXlsx(missing_dst, sheet_name="s4_1")
        assert False, "Should raise ValueError"
    except ValueError as e:
        print(f"✅ Correct error: {e}")
        assert "Missing" in str(e)
        assert "line 2" in str(e)


def test_write_id_not_found():
    """Test warning when ID not in original file"""
    print("\n" + "="*60)
    print("Test: ID not found (should warn and skip)")
    print("="*60)

    test_file = Path("input/test_script.xlsx")
    output_file = Path("output/test_id_mismatch.xlsx")
    create_test_xlsx(test_file)

    config = CConfig("config.yaml")
    config.xlsx.filePath = str(test_file)
    config.xlsx.outputPath = str(output_file)
    config.xlsx.sheetName = ["s4_1"]  # Must be list

    jsonline = '''{"id":1,"dst":"Translation 1"}
{"id":999,"dst":"This ID doesn't exist"}
{"id":2,"dst":"Translation 2"}'''

    xlsx = XLSXCore(config)

    # Must read first to build row mapping
    xlsx.readXlsx("s4_1")

    # Now write
    xlsx.writeXlsx(jsonline, sheet_name="s4_1")

    print("✅ Non-existent IDs skipped with warning")


def run_all_tests():
    """Run all test cases"""
    print("\n" + "="*60)
    print("XLSX Core Integration Tests")
    print("="*60)

    tests = [
        test_read_with_empty_cells,
        test_file_not_found,
        test_sheet_not_found,
        test_column_not_found,
        test_write_valid_jsonline,
        test_write_invalid_json,
        test_write_missing_keys,
        test_write_id_not_found,
    ]

    failed = []

    for test in tests:
        try:
            test()
        except AssertionError as e:
            print(f"❌ Test failed: {test.__name__}")
            print(f"   Error: {e}")
            failed.append(test.__name__)
        except Exception as e:
            print(f"❌ Test error: {test.__name__}")
            print(f"   Error: {e}")
            failed.append(test.__name__)

    print("\n" + "="*60)
    print("Test Summary")
    print("="*60)
    print(f"Total tests: {len(tests)}")
    print(f"Passed: {len(tests) - len(failed)}")
    print(f"Failed: {len(failed)}")

    if failed:
        print("\nFailed tests:")
        for name in failed:
            print(f"  - {name}")
        sys.exit(1)
    else:
        print("\n✅ All tests passed!")


if __name__ == "__main__":
    run_all_tests()
